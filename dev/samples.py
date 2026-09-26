#!/usr/bin/env python3
"""Example files for exercising the canvas viewers (developer mode only).
usage: python3 dev/samples.py <out-dir>   (the app calls this via window.__dev.samples())
Each generator is independent: a missing library skips that file instead of failing the rest."""
import csv, io, json, os, sys, zipfile

out = sys.argv[1] if len(sys.argv) > 1 else "workspace/uploads/samples"
os.makedirs(out, exist_ok=True)
made, skipped = [], []

def job(name):
    def wrap(fn):
        try:
            fn(os.path.join(out, name)); made.append(name)
        except Exception as e:  # noqa: BLE001
            skipped.append(f"{name}: {type(e).__name__}: {e}")
        return fn
    return wrap

ROWS = [("Month", "Revenue", "Costs"), ("Jan", 4200, 3100), ("Feb", 4650, 3240), ("Mar", 5100, 3300), ("Apr", 4980, 3420), ("May", 5600, 3510), ("Jun", 6120, 3600)]

@job("sales.csv")
def _(p):
    with open(p, "w", newline="") as f: csv.writer(f).writerows(ROWS)

@job("budget.xlsx")
def _(p):
    from openpyxl import Workbook
    wb = Workbook(); ws = wb.active; ws.title = "Summary"
    for r in ROWS: ws.append(r)
    ws.append(("Total", "=SUM(B2:B7)", "=SUM(C2:C7)"))
    ws2 = wb.create_sheet("Staff")
    for r in [("Name", "Role", "Start"), ("Asha", "Design", "2024-03-01"), ("Ravi", "Backend", "2023-11-15"), ("Mei", "Data", "2025-01-06")]: ws2.append(r)
    wb.save(p)

@job("report.docx")
def _(p):
    from docx import Document
    d = Document(); d.add_heading("Quarterly report", 0)
    d.add_paragraph("Revenue grew 46% from January to June while costs rose 16%.")
    d.add_heading("Highlights", 1)
    for t in ("Two new enterprise customers", "Support backlog halved", "Mobile release shipped"): d.add_paragraph(t, style="List Bullet")
    t = d.add_table(rows=1, cols=3); t.style = "Light Grid"
    for i, h in enumerate(ROWS[0]): t.rows[0].cells[i].text = h
    for r in ROWS[1:]:
        c = t.add_row().cells
        for i, v in enumerate(r): c[i].text = str(v)
    d.save(p)

@job("deck.pptx")
def _(p):
    from pptx import Presentation
    from pptx.util import Inches
    pr = Presentation()
    s = pr.slides.add_slide(pr.slide_layouts[0]); s.shapes.title.text = "Roadmap"; s.placeholders[1].text = "H2 planning"
    s = pr.slides.add_slide(pr.slide_layouts[1]); s.shapes.title.text = "Goals"
    tf = s.placeholders[1].text_frame; tf.text = "Ship offline mode"
    for t in ("Cut cold start to 1 s", "Self-serve billing"): tf.add_paragraph().text = t
    s = pr.slides.add_slide(pr.slide_layouts[5]); s.shapes.title.text = "Numbers"
    tb = s.shapes.add_table(len(ROWS), 3, Inches(1), Inches(1.6), Inches(8), Inches(3)).table
    for i, r in enumerate(ROWS):
        for j, v in enumerate(r): tb.cell(i, j).text = str(v)
    pr.save(p)

@job("chart.png")
def _(p):
    import matplotlib; matplotlib.use("Agg")
    import matplotlib.pyplot as plt
    fig, ax = plt.subplots(figsize=(6, 3.4), dpi=120)
    ax.plot([r[0] for r in ROWS[1:]], [r[1] for r in ROWS[1:]], label="Revenue"); ax.plot([r[0] for r in ROWS[1:]], [r[2] for r in ROWS[1:]], label="Costs")
    ax.legend(frameon=False); ax.spines[["top", "right"]].set_visible(False); fig.tight_layout(); fig.savefig(p)

@job("paper.pdf")
def _(p):
    import matplotlib; matplotlib.use("Agg")
    import matplotlib.pyplot as plt
    from matplotlib.backends.backend_pdf import PdfPages
    with PdfPages(p) as pdf:
        for i, (title, body) in enumerate([("A short paper", "Abstract. We measure how streaming renderers behave under bursty token arrival."), ("Method", "Tokens arrive in chunks of 1–60 characters with 5–160 ms gaps."), ("Results", "Adaptive pacing removes visible jitter without adding latency.")]):
            fig = plt.figure(figsize=(8.27, 11.69)); fig.text(0.1, 0.9, title, fontsize=22); fig.text(0.1, 0.84, body, fontsize=11, wrap=True); fig.text(0.5, 0.04, str(i + 1), ha="center")
            pdf.savefig(fig); plt.close(fig)

@job("notes.md")
def _(p):
    open(p, "w").write("# Notes\n\n- [x] Draft outline\n- [ ] Review numbers\n\n| Metric | Value |\n|---|---|\n| Users | 1,240 |\n| Churn | 2.1% |\n\n$$E = mc^2$$\n")

@job("config.json")
def _(p):
    json.dump({"name": "demo", "version": "1.2.0", "features": {"offline": True, "sync": False}}, open(p, "w"), indent=2)

@job("project.zip")
def _(p):
    with zipfile.ZipFile(p, "w", zipfile.ZIP_DEFLATED) as z:
        z.writestr("project/README.md", "# Project\nSample archive for the zip viewer.\n")
        z.writestr("project/src/main.py", "print('hello')\n")
        z.writestr("project/data/sales.csv", "\n".join(",".join(map(str, r)) for r in ROWS))
        for f in ("chart.png", "notes.md"):
            if os.path.exists(os.path.join(out, f)): z.write(os.path.join(out, f), "project/assets/" + f)

print(json.dumps({"dir": out, "made": made, "skipped": skipped}))
